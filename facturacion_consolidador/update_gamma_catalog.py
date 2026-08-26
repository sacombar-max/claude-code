import re
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

# Uso: python update_gamma_catalog.py Consolidador_Facturacion.xlsx catalogo_nuevo.xlsx
#
# Reemplaza el contenido de la pestaña Gamma_Catalogo con un catálogo nuevo (un Excel
# con columnas Ip Code, Description, R/C, Category, Mks description, en ese orden — como
# valores, no fórmulas) y ajusta automáticamente el rango de las fórmulas de Consolidado
# que buscan en Gamma_Catalogo, para que sigan apuntando a todas las filas del catálogo
# nuevo. Después de correr esto hay que recalcular (recalc.py).

WORKBOOK = sys.argv[1] if len(sys.argv) > 1 else "Consolidador_Facturacion.xlsx"
CATALOG = sys.argv[2] if len(sys.argv) > 2 else "catalogo_nuevo.xlsx"

BLUE_INPUT = Font(name="Arial", color="0000FF")
GAMMA_HEADER_ROW = 2

cat = pd.read_excel(CATALOG)
cat.columns = ["Ip Code", "Description", "R/C", "Category", "Mks description"]
cat = cat.dropna(subset=["Ip Code"]).drop_duplicates(subset=["Ip Code"]).sort_values("Ip Code")
cat["Ip Code"] = cat["Ip Code"].astype(int)

wb = load_workbook(WORKBOOK)
ws = wb["Gamma_Catalogo"]

# Limpia todo el contenido de datos existente (mas alla del nuevo catalogo tambien,
# por si el catalogo nuevo tiene menos filas que el anterior).
for r in range(GAMMA_HEADER_ROW + 1, ws.max_row + 1):
    for c in range(1, 6):
        ws.cell(row=r, column=c, value=None)

r = GAMMA_HEADER_ROW + 1
for _, row in cat.iterrows():
    ws.cell(row=r, column=1, value=int(row["Ip Code"])).font = BLUE_INPUT
    ws.cell(row=r, column=2, value=row["Description"]).font = BLUE_INPUT
    ws.cell(row=r, column=3, value=row["R/C"]).font = BLUE_INPUT
    ws.cell(row=r, column=4, value=row["Category"]).font = BLUE_INPUT
    ws.cell(row=r, column=5, value=row["Mks description"]).font = BLUE_INPUT
    r += 1
new_last_row = r - 1

# Ajusta los rangos Gamma_Catalogo!$X$3:$X$<N> (X en A,C,D,E) en todas las formulas de Consolidado.
pattern = re.compile(r"(Gamma_Catalogo!\$[A-E]\$3:\$[A-E]\$)(\d+)")
consolidado = wb["Consolidado"]
updated = 0
for row in consolidado.iter_rows():
    for cell in row:
        v = cell.value
        if isinstance(v, str) and v.startswith("=") and "Gamma_Catalogo" in v:
            new_v = pattern.sub(lambda m: f"{m.group(1)}{new_last_row}", v)
            if new_v != v:
                cell.value = new_v
                updated += 1

wb.save(WORKBOOK)
print(f"listo: {len(cat)} productos en Gamma_Catalogo (hasta fila {new_last_row}), "
      f"{updated} celdas de Consolidado con el rango actualizado")
