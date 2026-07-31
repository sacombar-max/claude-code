import sys

from openpyxl import load_workbook

# Uso: python hide_blank_rows.py Consolidador_Facturacion.xlsx
#
# Oculta (no borra) las filas de la pestaña Consolidado que no tienen datos todavía,
# para que al abrir el archivo solo se vean las filas con información. Requiere que
# el libro ya haya sido recalculado (recalc.py), para leer los valores de 'mercado'.
# Al mes siguiente, si hay más o menos filas con datos, hay que volver a correr esto.

WORKBOOK = sys.argv[1] if len(sys.argv) > 1 else "Consolidador_Facturacion.xlsx"

wb_values = load_workbook(WORKBOOK, data_only=True)
mercado_values = [c[0].value for c in wb_values["Consolidado"].iter_rows(min_row=2, min_col=1, max_col=1)]

wb = load_workbook(WORKBOOK)
ws = wb["Consolidado"]
hidden = 0
for i, value in enumerate(mercado_values, start=2):
    is_blank = value is None or value == ""
    ws.row_dimensions[i].hidden = is_blank
    hidden += 1 if is_blank else 0

wb.save(WORKBOOK)
print(f"listo: {hidden} filas en blanco ocultas en Consolidado de {WORKBOOK}")
