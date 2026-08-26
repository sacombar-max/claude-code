import sys
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Uso: python build_resumen.py Consolidador_Facturacion.xlsx
# Requiere que el libro ya tenga Brasil/Suiza_DE_CN_ID con datos y haya sido
# recalculado (recalc.py) al menos una vez, para poder leer los valores de Consolidado.
IN_OUT = sys.argv[1] if len(sys.argv) > 1 else "Consolidador_Facturacion.xlsx"

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF")
BOLD = Font(name=FONT_NAME, bold=True)
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


df_full = pd.read_excel(IN_OUT, sheet_name="Consolidado", header=0)
CONSOL_LAST_ROW = 1 + len(df_full)
df = df_full[df_full["cliente"].notna()]

mercados = sorted(df["mercado"].dropna().unique())
segmentos = sorted(df["Mks description"].dropna().unique())

wb = load_workbook(IN_OUT)
if "Resumen" in wb.sheetnames:
    del wb["Resumen"]
ws = wb.create_sheet("Resumen")

ws["A1"] = "Piezas facturadas por mercado y segmento (Mks description), y Amount total por mercado"
ws["A1"].font = Font(name=FONT_NAME, italic=True, color="7F7F7F")

HEADER_ROW = 3
ws.cell(row=HEADER_ROW, column=1, value="mercado")
for j, seg in enumerate(segmentos, start=2):
    ws.cell(row=HEADER_ROW, column=j, value=seg)
amount_col = len(segmentos) + 2
ws.cell(row=HEADER_ROW, column=amount_col, value="Amount Total")
style_header(ws, HEADER_ROW, amount_col)

def col_rng(col):
    return f"Consolidado!${col}$2:${col}${CONSOL_LAST_ROW}"

mercado_rng = col_rng("A")
qty_rng = col_rng("F")
amount_rng = col_rng("J")
mks_rng = col_rng("N")

r = HEADER_ROW + 1
for mercado in mercados:
    ws.cell(row=r, column=1, value=mercado)
    for j, seg in enumerate(segmentos, start=2):
        col_letter = get_column_letter(j)
        ws.cell(
            row=r, column=j,
            value=f'=SUMIFS({qty_rng},{mercado_rng},$A{r},{mks_rng},{col_letter}${HEADER_ROW})'
        )
    ws.cell(row=r, column=amount_col, value=f'=SUMIFS({amount_rng},{mercado_rng},$A{r})')
    r += 1
last_row = r - 1

total_row = r
ws.cell(row=total_row, column=1, value="Total").font = BOLD
for j in range(2, amount_col + 1):
    col_letter = get_column_letter(j)
    cell = ws.cell(row=total_row, column=j, value=f"=SUM({col_letter}{HEADER_ROW + 1}:{col_letter}{last_row})")
    cell.font = BOLD

widths = [16] + [14] * len(segmentos) + [14]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = f"B{HEADER_ROW + 1}"
ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(amount_col)}{last_row}"

wb.save(IN_OUT)
print("saved", IN_OUT, "mercados:", len(mercados), "segmentos:", len(segmentos))
