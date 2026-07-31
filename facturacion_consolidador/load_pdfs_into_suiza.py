import sys

from openpyxl import load_workbook
from openpyxl.styles import Font

from parse_proformas_pdf import collect_pdfs, parse_pdf

# Uso: python load_pdfs_into_suiza.py Consolidador_Facturacion.xlsx carpeta_con_pdfs/ [...]
#
# Lee proformas en PDF (una por cliente/pedido) y las escribe directo en la pestaña
# Suiza_DE_CN_ID del libro, reemplazando lo que hubiera ahí. Después de esto corre
# recalc.py y, si quieres, build_resumen.py.

SUIZA_CAP = 450  # debe coincidir con BRASIL_CAP/SUIZA_CAP de build_consolidador.py
ALIAS_FORMULA = (
    '=IF(A{r}="","",IFERROR(INDEX(Mercados!$A$2:$A$29,MATCH(A{r},Mercados!$B$2:$B$29,0)),"Revisar alias"))'
)


def main():
    args = sys.argv[1:]
    if len(args) < 2:
        print("Uso: python load_pdfs_into_suiza.py Consolidador_Facturacion.xlsx carpeta_con_pdfs/ [...]")
        sys.exit(1)

    workbook_path, *pdf_args = args
    pdf_paths = collect_pdfs(pdf_args)
    if not pdf_paths:
        print("No se encontraron PDFs en los argumentos dados.")
        sys.exit(1)

    all_rows = []
    for p in pdf_paths:
        all_rows.extend(parse_pdf(p))

    if len(all_rows) > SUIZA_CAP:
        print(f"ADVERTENCIA: {len(all_rows)} filas superan la capacidad actual ({SUIZA_CAP}). Avisa para ampliarla.")

    wb = load_workbook(workbook_path)
    ws = wb["Suiza_DE_CN_ID"]

    last_row = max(ws.max_row, 1 + SUIZA_CAP)
    for r in range(2, last_row + 1):
        for c in [1, 3, 4, 5, 6, 7, 8, 9, 10]:
            ws.cell(row=r, column=c, value=None)
        ws.cell(row=r, column=2, value=ALIAS_FORMULA.format(r=r))

    for i, row in enumerate(all_rows, start=2):
        ws.cell(row=i, column=1, value=row["cliente"])
        ws.cell(row=i, column=2, value=row["Tscode"])  # dato real del PDF, sobrescribe la fórmula de alias
        ws.cell(row=i, column=3, value=row["pf"])
        ws.cell(row=i, column=4, value=row["Ip Code"])
        ws.cell(row=i, column=5, value=row["Description"])
        ws.cell(row=i, column=6, value=row["Brand Line"])
        ws.cell(row=i, column=7, value=row["Quantity"])
        ws.cell(row=i, column=8, value=row["Net Price"])
        ws.cell(row=i, column=9, value=row["Amount"])
        ws.cell(row=i, column=10, value=row["Goods Origin"])

    wb.save(workbook_path)
    print(f"listo: {len(all_rows)} filas de {len(pdf_paths)} PDF(s) escritas en Suiza_DE_CN_ID de {workbook_path}")
    print("Ahora corre recalc.py (y build_resumen.py si quieres el resumen) sobre este archivo.")


if __name__ == "__main__":
    main()
