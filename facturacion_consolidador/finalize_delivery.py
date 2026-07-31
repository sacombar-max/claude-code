import sys

from openpyxl import load_workbook

# Uso: python finalize_delivery.py Consolidador_Facturacion.xlsx
#
# Deja el libro listo para entregar: borra la pestaña Instrucciones y oculta
# Mercados y Gamma_Catalogo (siguen ahí y las fórmulas las siguen usando,
# solo no se muestran). Visibles: Brasil, Suiza_DE_CN_ID, Consolidado, Resumen.

WORKBOOK = sys.argv[1] if len(sys.argv) > 1 else "Consolidador_Facturacion.xlsx"

wb = load_workbook(WORKBOOK)

if "Instrucciones" in wb.sheetnames:
    del wb["Instrucciones"]

for name in ["Mercados", "Gamma_Catalogo"]:
    if name in wb.sheetnames:
        wb[name].sheet_state = "hidden"

for name in ["Brasil", "Suiza_DE_CN_ID", "Consolidado", "Resumen"]:
    if name in wb.sheetnames:
        wb[name].sheet_state = "visible"

wb.save(WORKBOOK)
print(f"listo: {WORKBOOK} sin 'Instrucciones', con Mercados/Gamma_Catalogo ocultas")
